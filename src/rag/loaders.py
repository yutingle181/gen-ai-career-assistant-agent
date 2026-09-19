"""文档解析：PDF / DOCX / Markdown / TXT / HTML。

对应 JD 职责 3「文档解析」。统一输出 RawDocument(source, page, text)。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from ..logging_setup import get_logger
from .types import RawDocument

logger = get_logger(__name__)

SUPPORTED_SUFFIX = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"}


def _load_pdf(path: Path) -> list[RawDocument]:
    from pypdf import PdfReader

    docs: list[RawDocument] = []
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        if text.strip():
            docs.append(RawDocument(source=path.name, text=text, page=i))
    return docs


def _load_docx(path: Path) -> list[RawDocument]:
    import docx  # python-docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs if p.text)
    return [RawDocument(source=path.name, text=text, page=None)] if text.strip() else []


def _load_text(path: Path) -> list[RawDocument]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [RawDocument(source=path.name, text=text)] if text.strip() else []


def _load_html(path: Path) -> list[RawDocument]:
    from bs4 import BeautifulSoup

    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    return [RawDocument(source=path.name, text=text)] if text.strip() else []


def _file_time(path: Path) -> str:
    """文件最后修改时间（入库时写进 meta，供数据时效判断）。

    为什么要显式记录：此前片段时间只能靠「源文件还在不在」反推，
    文件被移走 / 清理后，时效就**永久变成「未知」**，答案也就再也无法声明资料时间。
    入库时记一次，此后与源文件是否还在无关。
    """
    try:
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    except OSError:
        return ""


def load_file(path: str | Path) -> list[RawDocument]:
    """按后缀分发解析器；解析结果会带上文件时间（`meta["updated_at"]`）。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_SUFFIX:
        logger.warning("不支持的文件类型：%s", p.name)
        return []
    try:
        if suffix == ".pdf":
            docs = _load_pdf(p)
        elif suffix == ".docx":
            docs = _load_docx(p)
        elif suffix in {".html", ".htm"}:
            docs = _load_html(p)
        else:
            docs = _load_text(p)
    except Exception as exc:  # noqa: BLE001
        logger.error("解析失败 %s：%s", p.name, exc)
        return []

    stamp = _file_time(p)
    if stamp:
        for doc in docs:
            doc.meta.setdefault("updated_at", stamp)
    return docs


def load_dir(directory: str | Path, recursive: bool = True) -> list[RawDocument]:
    """批量解析目录内所有支持的文件。"""
    base = Path(directory)
    if not base.exists():
        return []
    pattern = "**/*" if recursive else "*"
    docs: list[RawDocument] = []
    for p in sorted(base.glob(pattern)):
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIX:
            docs.extend(load_file(p))
    return docs


def load_paths(paths) -> list[RawDocument]:
    """解析一组路径（文件或目录混合）。"""
    docs: list[RawDocument] = []
    for item in paths or []:
        p = Path(item)
        if p.is_dir():
            docs.extend(load_dir(p))
        elif p.is_file():
            docs.extend(load_file(p))
    return docs
